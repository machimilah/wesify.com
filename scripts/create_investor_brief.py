from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "BO_Investor_Value_Proposition.docx"
WAVE = ROOT / "metallic_wave_webpage_images" / "varied_positions_and_angles" / "wave_10.png"

INK = "171A17"
GRAPHITE = "3B423F"
MUTED = "666E67"
LINE = "DCDED9"
SOFT = "F4F5F2"
PALE = "E9ECE7"
WHITE = "FFFFFF"
AMBER = "C99A3F"


def rgb(hex_value: str) -> RGBColor:
    return RGBColor.from_string(hex_value)


def set_cell_fill(cell, color: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def set_cell_border(cell, **edges):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge_name, values in edges.items():
        edge = borders.find(qn(f"w:{edge_name}"))
        if edge is None:
            edge = OxmlElement(f"w:{edge_name}")
            borders.append(edge)
        for key, value in values.items():
            edge.set(qn(f"w:{key}"), str(value))


def set_cell_margins(cell, top=120, start=140, bottom=120, end=140):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_widths(table, widths):
    table.autofit = False
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(int(width.inches * 1440)))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = widths[idx]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(int(widths[idx].inches * 1440)))
            tc_w.set(qn("w:type"), "dxa")


def shade_run(run, color=INK):
    run.font.color.rgb = rgb(color)


def set_run(run, size=10, bold=False, color=INK, name="Arial", italic=False):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = rgb(color)
    return run


def set_para(paragraph, before=0, after=6, line=1.05, keep=False):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    fmt.keep_with_next = keep
    return paragraph


def add_text(doc, text, size=10, color=INK, bold=False, before=0, after=6,
             align=None, keep=False, italic=False):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    set_para(p, before=before, after=after, line=1.08, keep=keep)
    set_run(p.add_run(text), size=size, bold=bold, color=color, italic=italic)
    return p


def add_kicker(doc, text):
    p = doc.add_paragraph()
    set_para(p, before=0, after=7, keep=True)
    run = set_run(p.add_run(text.upper()), size=8, bold=True, color=GRAPHITE)
    run.font.letter_spacing = Pt(1.2)
    return p


def add_title(doc, title, subtitle=None):
    p = doc.add_paragraph()
    set_para(p, before=0, after=8, line=0.94, keep=True)
    set_run(p.add_run(title), size=25, bold=True, color=INK)
    if subtitle:
        add_text(doc, subtitle, size=10.5, color=MUTED, after=13)
    return p


def add_section_title(doc, number, title, subtitle=None):
    add_kicker(doc, f"{number} / INVESTMENT CASE")
    add_title(doc, title, subtitle)


def add_rule(doc, color=LINE, after=10):
    p = doc.add_paragraph()
    set_para(p, before=0, after=after)
    p_pr = p._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    borders.append(bottom)
    p_pr.append(borders)
    return p


def add_callout(doc, label, headline, body, fill=SOFT):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_widths(table, [Inches(7.08)])
    cell = table.cell(0, 0)
    set_cell_fill(cell, fill)
    set_cell_margins(cell, top=190, start=220, bottom=180, end=220)
    set_cell_border(cell, top={"val": "single", "sz": "10", "color": GRAPHITE})
    p = cell.paragraphs[0]
    set_para(p, after=5)
    set_run(p.add_run(label.upper()), size=7.5, bold=True, color=MUTED)
    p = cell.add_paragraph()
    set_para(p, after=5, line=1.0)
    set_run(p.add_run(headline), size=14, bold=True, color=INK)
    p = cell.add_paragraph()
    set_para(p, after=0, line=1.08)
    set_run(p.add_run(body), size=9.5, color=GRAPHITE)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_stat_row(doc, stats):
    table = doc.add_table(rows=1, cols=len(stats))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    width = 7.08 / len(stats)
    set_table_widths(table, [Inches(width)] * len(stats))
    for idx, (value, label) in enumerate(stats):
        cell = table.cell(0, idx)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_fill(cell, SOFT)
        set_cell_margins(cell, top=170, start=160, bottom=160, end=160)
        if idx:
            set_cell_border(cell, start={"val": "single", "sz": "5", "color": LINE})
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_para(p, after=3)
        set_run(p.add_run(value), size=19, bold=True, color=INK)
        p = cell.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_para(p, after=0, line=1.0)
        set_run(p.add_run(label), size=8, color=MUTED)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_cards(doc, cards, columns=3, fill=SOFT, title_size=10.5, body_size=8.5):
    rows = (len(cards) + columns - 1) // columns
    table = doc.add_table(rows=rows, cols=columns)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    gap = 0.08
    width = (7.08 - gap * (columns - 1)) / columns
    set_table_widths(table, [Inches(width)] * columns)
    for index in range(rows * columns):
        cell = table.cell(index // columns, index % columns)
        set_cell_margins(cell, top=150, start=160, bottom=145, end=160)
        set_cell_fill(cell, fill if index < len(cards) else WHITE)
        if index >= len(cards):
            continue
        label, title, body = cards[index]
        p = cell.paragraphs[0]
        set_para(p, after=5)
        set_run(p.add_run(label.upper()), size=7.2, bold=True, color=MUTED)
        p = cell.add_paragraph()
        set_para(p, after=5, line=1.0)
        set_run(p.add_run(title), size=title_size, bold=True, color=INK)
        p = cell.add_paragraph()
        set_para(p, after=0, line=1.06)
        set_run(p.add_run(body), size=body_size, color=GRAPHITE)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_two_column_rows(doc, rows, left_width=2.18, right_width=4.9, header=None):
    table = doc.add_table(rows=1 if header else 0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_widths(table, [Inches(left_width), Inches(right_width)])
    if header:
        for idx, text in enumerate(header):
            cell = table.cell(0, idx)
            set_cell_fill(cell, GRAPHITE)
            set_cell_margins(cell, top=120, start=150, bottom=120, end=150)
            p = cell.paragraphs[0]
            set_para(p, after=0)
            set_run(p.add_run(text.upper()), size=7.5, bold=True, color=WHITE)
    for label, body in rows:
        cells = table.add_row().cells
        for cell in cells:
            set_cell_margins(cell, top=130, start=150, bottom=125, end=150)
            set_cell_border(cell, bottom={"val": "single", "sz": "4", "color": LINE})
        p = cells[0].paragraphs[0]
        set_para(p, after=0)
        set_run(p.add_run(label), size=9, bold=True, color=INK)
        p = cells[1].paragraphs[0]
        set_para(p, after=0, line=1.05)
        set_run(p.add_run(body), size=8.7, color=GRAPHITE)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_bullet(doc, title, body, size=9.3):
    p = doc.add_paragraph(style="List Bullet")
    set_para(p, before=0, after=7, line=1.06)
    p.paragraph_format.left_indent = Inches(0.22)
    p.paragraph_format.first_line_indent = Inches(-0.16)
    set_run(p.add_run(title + " "), size=size, bold=True, color=INK)
    set_run(p.add_run(body), size=size, color=GRAPHITE)
    return p


def add_page_break(doc):
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)


def add_page_field(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for element in (begin, instruction, separate, text, end):
        run._r.append(element)
    set_run(run, size=8, color=MUTED)


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.62)
    section.bottom_margin = Inches(0.58)
    section.left_margin = Inches(0.71)
    section.right_margin = Inches(0.71)
    section.header_distance = Inches(0.25)
    section.footer_distance = Inches(0.25)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    normal.font.size = Pt(10)
    normal.font.color.rgb = rgb(INK)

    for name, size, bold, color in (
        ("Title", 25, True, INK),
        ("Subtitle", 11, False, MUTED),
        ("Heading 1", 18, True, INK),
        ("Heading 2", 12, True, INK),
        ("Heading 3", 10, True, GRAPHITE),
    ):
        style = styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.color.rgb = rgb(color)

    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_para(p, after=2)
    set_run(p.add_run("BO"), size=8, bold=True, color=INK)
    set_run(p.add_run("   |   INVESTOR VALUE PROPOSITION"), size=7.5, color=MUTED)
    add_rule_to_paragraph(p)

    footer = section.footer
    table = footer.add_table(rows=1, cols=2, width=Inches(7.08))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_widths(table, [Inches(5.7), Inches(1.38)])
    left, right = table.rows[0].cells
    for cell in (left, right):
        set_cell_margins(cell, top=0, start=0, bottom=0, end=0)
    p = left.paragraphs[0]
    set_para(p, after=0)
    set_run(p.add_run("PREPARED FOR INVESTOR CONVERSATIONS  |  AUGUST 2026"), size=7, color=MUTED)
    add_page_field(right.paragraphs[0])


def add_rule_to_paragraph(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "4")
    bottom.set(qn("w:space"), "5")
    bottom.set(qn("w:color"), LINE)
    borders.append(bottom)
    p_pr.append(borders)


def cover(doc):
    add_text(doc, "BO", size=11, bold=True, color=INK, after=8)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_para(p, after=18)
    p.add_run().add_picture(str(WAVE), width=Inches(7.08), height=Inches(3.38))

    add_kicker(doc, "Investor value proposition")
    p = doc.add_paragraph()
    set_para(p, before=0, after=12, line=0.9)
    set_run(p.add_run("The business operating system\nthat builds itself around the company."), size=27, bold=True, color=INK)
    add_text(
        doc,
        "BO turns a plain-language description into a tailored Command Center: the records, workflows, metrics, controls, and connected data a company needs to operate.",
        size=11,
        color=GRAPHITE,
        after=18,
    )
    add_cards(doc, [
        ("01", "Immediate value", "A working system appears before the buyer commits to a configuration project."),
        ("02", "Structural differentiation", "The product compiles from company evidence instead of exposing a generic suite."),
        ("03", "Compounding moat", "Anonymous industry-level corrections make later builds more accurate."),
    ], columns=3, fill=SOFT, title_size=10, body_size=8.2)


def page_problem(doc):
    add_section_title(doc, "01", "A large gap between spreadsheets and ERP", "Smaller operating companies are forced to choose between fragmentation and an implementation project.")
    add_cards(doc, [
        ("TODAY", "Fragmented tools", "Payments, CRM, accounting, scheduling, and spreadsheets each hold part of the truth. Teams re-key data and reconcile by hand."),
        ("ALTERNATIVE", "Overbuilt suites", "ERP and CRM suites ship broad module catalogs, then transfer the burden of configuration, training, and process design to the buyer."),
        ("CONSEQUENCE", "Operational blind spots", "Ownership, handoffs, exceptions, and definitions live in people rather than in one usable operating model."),
    ], columns=3, fill=SOFT)
    add_callout(
        doc,
        "BO's wedge",
        "Outgrown spreadsheets. Not ready for a six-month ERP program.",
        "The initial target is an operationally real company, roughly 5-200 employees, without a dedicated systems administrator. The strongest early verticals have specific workflows: field services, trades, logistics, studios, clinics, and small manufacturers.",
    )
    add_text(doc, "The category shift", size=12, bold=True, after=6, keep=True)
    add_two_column_rows(doc, [
        ("Traditional software", "Choose a product, select modules, configure fields, map processes, migrate data, train users."),
        ("BO", "Describe the business, answer only material gaps, review BO's understanding, receive the configured Command Center."),
        ("Investor implication", "The demonstration is the product. BO can prove relevance before asking a company to switch or pay."),
    ])
    add_text(doc, "No unsupported market-size or traction figures are used in this brief. The investment case rests on product architecture, customer economics, and milestones that can be measured.", size=8, color=MUTED, italic=True, before=8, after=0)


def page_product(doc):
    add_section_title(doc, "02", "One sentence becomes an operating model", "BO asks only what can change the result, then compiles company knowledge into software.")
    add_cards(doc, [
        ("1", "Describe", "The operator explains the company in ordinary language."),
        ("2", "Classify", "BO maps the company to an industry and operating archetype."),
        ("3", "Resolve", "High-information questions settle undecided capabilities."),
        ("4", "Compile", "Evidence becomes records, pages, KPIs, workflows, and controls."),
        ("5", "Connect", "Existing systems remain the source of record where appropriate."),
        ("6", "Evolve", "Approved changes create a new, versioned Command Center."),
    ], columns=3, fill=SOFT, title_size=10.5, body_size=8.1)
    add_callout(
        doc,
        "Product rule",
        "BO installs structure, never fiction.",
        "A generated workspace can contain record types, relationships, workflows, responsibilities, controls, and unconfigured metrics. It must not invent customers, transactions, employees, performance figures, or operational alerts.",
        fill=PALE,
    )
    add_text(doc, "A deeper operating model, without changing how BO works", size=12, bold=True, after=7, keep=True)
    add_two_column_rows(doc, [
        ("Responsibility", "RACI-style ownership and approvals become role and policy primitives."),
        ("Flow", "Inputs, outputs, handoffs, exceptions, and stage gates become process definitions."),
        ("Data", "Master data dictionaries become generated schemas with governed field definitions."),
        ("Performance", "KPI catalogs connect measures to source capabilities, owners, and decisions."),
        ("Automation", "Reusable automation patterns are selected only when company evidence supports them."),
    ], header=("Operating knowledge", "How BO turns it into product"))


def page_defensibility(doc):
    add_section_title(doc, "03", "A product that compounds by industry", "The catalog can be copied. The accumulated record of what real companies kept, removed, and added cannot.")
    add_stat_row(doc, [
        ("1,923", "industry titles"),
        ("120", "buildable capabilities"),
        ("25", "operating archetypes"),
        ("5+", "companies before evidence applies"),
    ])
    add_text(doc, "The learning loop", size=12, bold=True, after=7, keep=True)
    add_cards(doc, [
        ("BUILD", "BO proposes", "Rules, public research, and the operator's statements produce the first Command Center."),
        ("OPERATE", "The company corrects", "Kept, removed, and newly added capabilities become clean operational labels."),
        ("AGGREGATE", "Only counts survive", "No company name, workspace identifier, record, prompt, or field content enters shared evidence."),
        ("IMPROVE", "The next build starts better", "Observed behavior outranks industry research; explicit company statements still outrank both."),
    ], columns=2, fill=SOFT, title_size=10.5, body_size=8.5)
    add_text(doc, "Guardrails make the loop credible", size=12, bold=True, before=4, after=5, keep=True)
    add_two_column_rows(doc, [
        ("Minimum sample", "At least five companies contribute before a pattern can influence another company."),
        ("Strong threshold", "A 60% decision threshold avoids treating a narrow majority as an industry truth."),
        ("One company, one vote", "Repeated user activity cannot outweigh other companies."),
        ("Explicit beats inferred", "What this operator states always overrides research or observed patterns."),
    ])
    add_text(doc, "Cold-start reality: the mechanism is implemented, but observed evidence must be earned through adoption. Until then, BO relies on authored rules and cited industry research.", size=8, color=MUTED, italic=True, before=7, after=0)


def page_value(doc):
    add_section_title(doc, "04", "Value is created at setup and during operation", "BO reduces the cost of getting the right system and the cost of keeping it aligned as the company changes.")
    add_cards(doc, [
        ("OWNER / FOUNDER", "A company view, not an app view", "Customers, delivery, cash, responsibilities, and exceptions are visible through one operating model."),
        ("OPERATIONS LEAD", "Work has owners and states", "Processes expose handoffs, deadlines, approvals, and gaps instead of hiding them in chat and spreadsheets."),
        ("TEAM", "Less software interpretation", "Navigation and terminology follow the configured company rather than a vendor's universal object model."),
    ], columns=3, fill=SOFT)
    add_text(doc, "Economic mechanisms", size=12, bold=True, before=5, after=7, keep=True)
    add_two_column_rows(doc, [
        ("Lower implementation burden", "Discovery and compilation replace much of the blank-page configuration work."),
        ("Fewer duplicate systems", "A capability can be built in BO or connected to the system that already owns it."),
        ("Faster operational changes", "Natural-language requests become typed, reviewable change proposals with impact and rollback."),
        ("Improved control", "KPI definitions, responsibilities, data sources, and approval policies remain connected."),
        ("AI readiness", "Structured processes and trustworthy data create the foundation for governed agents and decision support."),
    ], header=("Value lever", "Why it matters"))
    add_callout(
        doc,
        "Core promise",
        "You prompt, I build",
        "BO is not trying to recreate every specialist system. It is the operating layer that decides what belongs in the company, builds what is missing, and connects what should remain elsewhere.",
    )


def page_business_model(doc):
    add_section_title(doc, "05", "A low-friction entry with expansion tied to value", "The personalized first build is free. Paid plans monetize scale, change, connections, and collaboration.")
    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_widths(table, [Inches(2.05), Inches(1.5), Inches(1.72), Inches(1.81)])
    headers = ["PLAN", "FREE", "PRO", "BUSINESS"]
    for idx, text in enumerate(headers):
        cell = table.cell(0, idx)
        set_cell_fill(cell, GRAPHITE if idx else INK)
        set_cell_margins(cell, top=130, start=130, bottom=125, end=130)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if idx else WD_ALIGN_PARAGRAPH.LEFT
        set_para(p, after=0)
        set_run(p.add_run(text), size=8, bold=True, color=WHITE)
    rows = [
        ("Price / month", "$0", "$10", "$50"),
        ("Workspaces", "1", "1", "5"),
        ("Records", "200", "Unlimited", "Unlimited"),
        ("Rebuilds / month", "0", "5", "20"),
        ("Connected apps", "-", "Included", "Included"),
        ("Team access", "-", "-", "Included"),
    ]
    for row_idx, row in enumerate(rows):
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            set_cell_fill(cells[idx], SOFT if row_idx % 2 == 0 else WHITE)
            set_cell_margins(cells[idx], top=115, start=130, bottom=110, end=130)
            set_cell_border(cells[idx], bottom={"val": "single", "sz": "4", "color": LINE})
            p = cells[idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if idx else WD_ALIGN_PARAGRAPH.LEFT
            set_para(p, after=0)
            set_run(p.add_run(value), size=8.7, bold=(idx == 0), color=INK if idx == 0 else GRAPHITE)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    add_text(doc, "Why the economics can improve with scale", size=12, bold=True, before=7, after=6, keep=True)
    add_bullet(doc, "First build as acquisition.", "A buyer sees a system shaped around their own company before paying.")
    add_bullet(doc, "Usage is structurally light.", "Once built, the Command Center is primarily persistent structured data and application runtime.")
    add_bullet(doc, "Research amortizes by industry.", "Industry-level research can be reused by every later company in that segment.")
    add_bullet(doc, "Vertical concentration improves both product and margin.", "More customers in one industry create stronger evidence and spread fixed knowledge costs.")
    add_callout(
        doc,
        "Commercial discipline",
        "No deletion as a collection tactic.",
        "A lapsed plan falls back to Free limits while existing data stays readable and exportable. The product sells ongoing capacity and capability, not fear of losing company records.",
        fill=PALE,
    )


def page_gtm(doc):
    add_section_title(doc, "06", "Win vertically, expand horizontally", "BO can address many industries, but early distribution should deliberately concentrate evidence and references.")
    add_cards(doc, [
        ("LAND", "Personalized free build", "The product creates its own demo from the company's description."),
        ("FOCUS", "A few operating verticals", "Concentration accelerates accuracy, references, and repeatable onboarding."),
        ("EXPAND", "Connections and team use", "Paid value grows as BO becomes the shared view across existing systems."),
        ("MOVE UP", "Trust and governance", "Roles, auditability, data controls, and reliable runtime support larger deployments."),
    ], columns=2, fill=SOFT)
    add_text(doc, "The next commercial proof", size=12, bold=True, before=5, after=6, keep=True)
    add_two_column_rows(doc, [
        ("Design partners", "Recruit five companies from one narrowly defined operating segment."),
        ("Recognition", "More than 80% of the central process is recognized as correct by operators."),
        ("Correction rate", "Fewer than 15% of important claims require correction."),
        ("Time to value", "A useful reviewed operating model is reached in under 20 minutes."),
        ("Operational outcome", "Each company identifies at least one concrete decision or control improved by BO."),
    ], header=("Validation gate", "Evidence to collect"))
    add_text(doc, "Metrics that reveal compounding", size=12, bold=True, before=7, after=5, keep=True)
    add_stat_row(doc, [
        ("↓", "corrections per build"),
        ("↑", "build-to-first-record"),
        ("↑", "connected apps / workspace"),
        ("↑", "free-to-paid by vertical"),
    ])


def page_state(doc):
    add_section_title(doc, "07", "The platform exists; adoption is the next risk to retire", "BO has moved beyond a visual prototype, but it should not claim commercial validation it has not earned.")
    add_two_column_rows(doc, [
        ("Built and tested", "Structured discovery; 1,923-title taxonomy; 120-capability catalog; dependency graph; generated schemas and Command Center; persistent records and builds; tenant access; accounts; billing; observability; deployment; versioning."),
        ("Connected data", "A real, read-only Stripe connector imports customers, subscriptions, and payments; credentials are encrypted and never returned; remote deletions are marked rather than silently applied."),
        ("Compounding layer", "Industry research with citations plus anonymous, thresholded capability and reusable-pattern learning."),
        ("Still partial", "Durable workflow runtime, granular server-enforced agent policies, broader connector lifecycle, unified human approval queues, product/engineering/security/legal/data-governance domain depth."),
        ("Not yet proven", "Repeatable acquisition, retention, willingness to pay, vertical unit economics, and accuracy improvement from real customer behavior."),
    ], header=("State", "Investor interpretation"))
    add_text(doc, "Capital should unlock evidence, not more surface area", size=12, bold=True, before=8, after=7, keep=True)
    add_cards(doc, [
        ("1", "Design-partner deployment", "Put BO into one vertical and instrument activation, correction, usage, and conversion."),
        ("2", "Durable execution", "Ship workflow instances, retries, exception handling, and a unified approval queue."),
        ("3", "Connector depth", "Add the few systems that dominate the chosen vertical, with monitored synchronization."),
        ("4", "Governed intelligence", "Enforce agent tool/data scopes, capture traces, and gate autonomy through evaluations."),
    ], columns=2, fill=SOFT, title_size=10.5, body_size=8.4)
    add_callout(
        doc,
        "Investment thesis",
        "BO can become the adaptive operating layer for companies underserved by both point tools and ERP suites.",
        "The near-term case is not 'AI for every business.' It is measurable: prove that a generated Command Center is recognized faster, corrected less, used earlier, and retained more deeply as industry evidence and connections accumulate.",
        fill=PALE,
    )


def page_risks(doc):
    add_section_title(doc, "08", "The opportunity is large because the hard parts are real", "A credible investor case states the risks and the mechanisms BO uses to contain them.")
    add_two_column_rows(doc, [
        ("Cold start", "Observed evidence begins empty. Mitigation: cited industry research, conservative authored rules, and vertical concentration."),
        ("Wrong inference", "A confident but incorrect workspace destroys trust. Mitigation: explicit evidence basis, playback, confidence, corrections, and stated answers overriding all inference."),
        ("System-of-record trust", "Operators are cautious with core data. Mitigation: tenant isolation, encrypted credentials, read-first connectors, versioning, audit, and no silent deletion."),
        ("Connector maintenance", "Every external API creates ongoing cost. Mitigation: a uniform provider contract and a narrow, vertical-first connector strategy."),
        ("Incumbent response", "Suites can add conversational setup. BO's defense is not chat; it is a narrow compiler, company evidence, and an industry learning loop that improves from operation."),
        ("Premature autonomy", "AI actions can create operational harm. Mitigation: deterministic validation, typed proposals, human approval, scoped tools, and progressive autonomy."),
    ], header=("Risk", "Response"))
    add_callout(
        doc,
        "In one sentence",
        "BO decides what a company needs instead of asking it to configure what it does not - and it is designed to decide better each time another company in the same industry uses it.",
        "That combines a strong product demonstration, a software-margin subscription model, and a privacy-safe data advantage that can compound with focused adoption.",
    )
    add_text(doc, "Source basis", size=12, bold=True, before=8, after=6, keep=True)
    add_text(
        doc,
        "Prepared from BO's live codebase and internal product materials, including the implementation architecture, roadmap, knowledge base, connected-app design, business research engine, billing model, and Oficial.docm. Product-state claims are limited to behavior represented in the current implementation and tests.",
        size=8.7,
        color=GRAPHITE,
        after=8,
    )
    add_text(doc, "Excluded by design: unsupported traction, market-size, ROI, and payroll/labor calculation claims.", size=8, color=MUTED, italic=True, after=0)


def build():
    doc = Document()
    configure_document(doc)
    core = doc.core_properties
    core.title = "BO Investor Value Proposition"
    core.subject = "Investor brief for BO, an AI-native Business Command Center"
    core.author = "BO"
    core.keywords = "BO, investor, Business Command Center, AI, operating system"
    core.comments = "Prepared from the current BO product and internal documentation."

    cover(doc)
    for page in (page_problem, page_product, page_defensibility, page_value, page_business_model, page_gtm, page_state, page_risks):
        add_page_break(doc)
        page(doc)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
